import base64
import io
import logging

from odoo import _, api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

TASK_HEADERS = [
    'name', 'description', 'target_model', 'verification_method',
    'verification_domain', 'verification_code', 'verification_count',
    'difficulty', 'required_modules', 'groups',
]
GROUP_HEADERS = ['name', 'icon', 'sequence']

VALID_METHODS = ('domain', 'python')
VALID_DIFFICULTIES = ('easy', 'medium', 'hard')


class SpeedrunTaskImportWizard(models.TransientModel):
    _name = 'speedrun.task.import.wizard'
    _description = 'Speedrun Task Import Wizard'

    import_file = fields.Binary(string='File', required=True)
    import_filename = fields.Char(string='File Name')
    auto_install_modules = fields.Boolean(
        string='Install required modules', default=True,
        help="Automatically install the Odoo modules referenced by the imported "
             "tasks (requires Administration / Settings rights).")
    update_existing = fields.Boolean(
        string='Update existing', default=False,
        help="If a group with the same name already exists, update it instead of "
             "reusing it as-is.")
    existing_group_ids = fields.Many2many(
        'speedrun.task.group', string='Existing groups', readonly=True,
        default=lambda self: self.env['speedrun.task.group'].search([]).ids)
    result_info = fields.Text(readonly=True)

    # ------------------------------------------------------------------
    def action_download_template(self):
        return {
            'type': 'ir.actions.act_url',
            'url': '/odoo_speedrun/task_import_template',
            'target': 'self',
        }

    # ------------------------------------------------------------------
    def _load_workbook(self):
        try:
            import openpyxl
        except ImportError:
            raise UserError(_("The 'openpyxl' library is required to import xlsx files."))
        if not self.import_file:
            raise UserError(_("Please upload a file first."))
        try:
            data = base64.b64decode(self.import_file)
            return openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        except Exception as exc:
            raise UserError(_("Could not read the file. Make sure it is a valid .xlsx "
                              "workbook (%s).", exc))

    @staticmethod
    def _sheet_rows(sheet, expected_headers):
        """Yield dicts mapping header -> value for each non-empty data row."""
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            return
        headers = [str(h).strip().lower() if h is not None else '' for h in header_row]
        for row in rows:
            if row is None or all(c is None or str(c).strip() == '' for c in row):
                continue
            record = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                record[header] = row[idx] if idx < len(row) else None
            yield record

    @staticmethod
    def _cell_str(value):
        if value is None:
            return ''
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    # ------------------------------------------------------------------
    def _import_groups(self, wb):
        Group = self.env['speedrun.task.group']
        cache = {g.name.strip().lower(): g for g in Group.search([])}
        created = 0
        if 'Groups' not in wb.sheetnames:
            return cache, created
        for row in self._sheet_rows(wb['Groups'], GROUP_HEADERS):
            name = self._cell_str(row.get('name'))
            if not name:
                continue
            icon = self._cell_str(row.get('icon'))
            seq_raw = self._cell_str(row.get('sequence'))
            vals = {'name': name}
            if icon:
                vals['icon'] = icon
            if seq_raw:
                try:
                    vals['sequence'] = int(float(seq_raw))
                except ValueError:
                    pass
            existing = cache.get(name.lower())
            if existing:
                if self.update_existing:
                    existing.write(vals)
            else:
                grp = Group.create(vals)
                cache[name.lower()] = grp
                created += 1
        return cache, created

    def _import_tasks(self, wb, group_cache):
        Task = self.env['speedrun.task']
        Group = self.env['speedrun.task.group']
        if 'Tasks' not in wb.sheetnames:
            raise UserError(_("The workbook must contain a 'Tasks' sheet."))
        created = 0
        module_names = set()
        for row in self._sheet_rows(wb['Tasks'], TASK_HEADERS):
            name = self._cell_str(row.get('name'))
            target = self._cell_str(row.get('target_model'))
            if not name or not target:
                continue
            method = self._cell_str(row.get('verification_method')).lower() or 'domain'
            if method not in VALID_METHODS:
                method = 'domain'
            difficulty = self._cell_str(row.get('difficulty')).lower() or 'medium'
            if difficulty not in VALID_DIFFICULTIES:
                difficulty = 'medium'
            count_raw = self._cell_str(row.get('verification_count'))
            try:
                count = int(float(count_raw)) if count_raw else 1
            except ValueError:
                count = 1
            required = self._cell_str(row.get('required_modules'))
            vals = {
                'name': name,
                'description': self._cell_str(row.get('description')),
                'target_model': target,
                'verification_method': method,
                'verification_domain': self._cell_str(row.get('verification_domain')) or '[]',
                'verification_code': self._cell_str(row.get('verification_code')),
                'verification_count': count,
                'difficulty': difficulty,
                'required_modules': required,
            }
            # Resolve groups by name (create on the fly if unknown)
            groups = self.env['speedrun.task.group']
            for gname in self._cell_str(row.get('groups')).split(','):
                gname = gname.strip()
                if not gname:
                    continue
                grp = group_cache.get(gname.lower())
                if not grp:
                    grp = Group.create({'name': gname})
                    group_cache[gname.lower()] = grp
                groups |= grp
            if groups:
                vals['group_ids'] = [(6, 0, groups.ids)]
            Task.create(vals)
            created += 1
            for mod in required.split(','):
                if mod.strip():
                    module_names.add(mod.strip())
        return created, module_names

    # ------------------------------------------------------------------
    def action_import(self):
        self.ensure_one()
        wb = self._load_workbook()
        group_cache, groups_created = self._import_groups(wb)
        tasks_created, module_names = self._import_tasks(wb, group_cache)

        if not tasks_created and not groups_created:
            raise UserError(_("Nothing to import: no valid rows found in the file."))

        _logger.info("Speedrun import: %s group(s), %s task(s) created by %s",
                     groups_created, tasks_created, self.env.user.login)

        # Install required modules if requested and allowed
        to_install = self.env['ir.module.module']
        if module_names:
            to_install = self.env['ir.module.module'].sudo().search([
                ('name', 'in', list(module_names)),
                ('state', '=', 'uninstalled'),
            ])
        if to_install and self.auto_install_modules and self.env.user.has_group('base.group_system'):
            _logger.info("Speedrun import: installing modules %s", to_install.mapped('name'))
            return to_install.button_immediate_install()

        # No install → notify and close
        msg = _("%(groups)s group(s) and %(tasks)s task(s) imported.",
                groups=groups_created, tasks=tasks_created)
        if to_install:
            msg += _(" Modules to install manually: %s.", ', '.join(to_install.mapped('name')))
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _("Import complete"),
                'message': msg,
                'type': 'success',
                'sticky': bool(to_install),
                'next': {'type': 'ir.actions.act_window_close'},
            },
        }
